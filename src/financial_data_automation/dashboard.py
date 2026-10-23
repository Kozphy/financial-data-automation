"""Weekly cash workbook. Actuals and hypotheticals stay on separate sheets."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "week_starting",
    "qualified_leads",
    "proposals_submitted",
    "responses",
    "response_rate",
    "discovery_calls",
    "paid_customers",
    "gross_revenue_usd",
    "platform_fees_usd",
    "refunds_usd",
    "operating_costs_usd",
    "net_cash_usd",
    "gumroad_purchases",
    "saas_paying_subscribers",
    "mrr_usd",
]


def write_dashboard(path: Path) -> None:
    workbook = Workbook()
    actuals = workbook.active
    actuals.title = "Actuals"
    actuals["A1"] = "Recorded transactions only. A blank or zero is an absence of evidence, not a forecast."
    actuals["A1"].font = Font(bold=True)
    for index, header in enumerate(HEADERS, start=1):
        cell = actuals.cell(3, index, header)
        cell.font = Font(bold=True)
        actuals.column_dimensions[get_column_letter(index)].width = max(len(header) + 2, 16)
    for row in range(4, 17):
        actuals.cell(row, 5, f'=IF(C{row}=0,"",D{row}/C{row})')
        actuals.cell(row, 12, f"=H{row}-I{row}-J{row}-K{row}")
        actuals.cell(row, 5).number_format = "0.0%"
        for column in (8, 9, 10, 11, 12, 15):
            actuals.cell(row, column).number_format = "#,##0.00"
    actuals["A4"] = "2026-10-23"
    for column in range(2, 16):
        if column not in (5, 12):
            actuals.cell(4, column, 0)

    forecast = workbook.create_sheet("Hypothetical")
    forecast["A1"] = "Hypothetical planning numbers. Do not add this sheet into Actuals."
    forecast["A1"].font = Font(bold=True)
    forecast["A3"] = "scenario"
    forecast["B3"] = "week"
    forecast["C3"] = "assumed_paid_customers"
    forecast["D3"] = "assumed_gross_usd"
    forecast["E3"] = "label"
    forecast["A4"] = "first_upwork_starter"
    forecast["B4"] = "2026-10-23"
    forecast["C4"] = 1
    forecast["D4"] = 99
    forecast["E4"] = "Not a recorded transaction."
    workbook.save(path)
