"""Formatted workbook contract for the client-facing summary."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from financial_data_automation.reporting import DISCLAIMER, write_reports


def test_summary_workbook_is_formatted_and_states_its_limit(tmp_path: Path) -> None:
    cleaned = pd.DataFrame(
        {
            "transaction_id": ["T1"],
            "date": [pd.Timestamp("2026-01-01")],
            "description": ["alpha"],
            "amount": [1250.5],
        }
    )
    summary = {
        "record_count": 1,
        "total_amount": 1250.5,
        "average_amount": 1250.5,
        "minimum_amount": 1250.5,
        "maximum_amount": 1250.5,
    }
    monthly = pd.DataFrame(
        {"month": ["2026-01"], "total_amount": [1250.5], "record_count": [1]}
    )
    write_reports(
        cleaned_df=cleaned,
        rejected_df=cleaned.iloc[0:0].copy(),
        summary=summary,
        monthly_df=monthly,
        quality_rows=[{"check": "clean_rows", "value": 1}],
        output_dir=tmp_path,
    )

    workbook = load_workbook(tmp_path / "financial_summary.xlsx")
    notice = workbook["notice"]
    assert notice["A1"].value == DISCLAIMER
    assert "not an audit opinion" in DISCLAIMER
    assert "forecast" in DISCLAIMER

    summary_sheet = workbook["summary"]
    assert summary_sheet.freeze_panes == "A2"
    assert summary_sheet["A1"].font.bold is True
    headers = [cell.value for cell in summary_sheet[1]]
    amount_col = headers.index("total_amount") + 1
    assert summary_sheet.cell(2, amount_col).number_format == "#,##0.00"

    monthly_sheet = workbook["monthly"]
    monthly_headers = [cell.value for cell in monthly_sheet[1]]
    monthly_amount = monthly_headers.index("total_amount") + 1
    assert monthly_sheet.cell(2, monthly_amount).number_format == "#,##0.00"
